import sys
import os
import json
import subprocess
import docx

def sanitize_str(val):
    if val is None:
        return ""
    s = str(val).strip()
    if s.lower() in ["undefined", "null", "unknown", "n/a", "0 kw"]:
        return ""
    return s

def format_currency(val):
    if val is None:
        return "£0.00"
    try:
        val_float = float(val)
        if val_float < 0:
            return f"-£{abs(val_float):,.2f}"
        return f"£{val_float:,.2f}"
    except (ValueError, TypeError):
        return "£0.00"

def replace_text_in_paragraph(p, search_str, replace_str):
    if search_str in p.text:
        for run in p.runs:
            if search_str in run.text:
                run.text = run.text.replace(search_str, replace_str)
                return
        full_text = p.text.replace(search_str, replace_str)
        p.runs[0].text = full_text
        for run in p.runs[1:]:
            run.text = ""

def set_cell_text(cell, text, align="left", bold=False, font_size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = docx.shared.Pt(2)
    p.paragraph_format.space_after = docx.shared.Pt(2)
    if align == "right":
        p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "center":
        p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.font.name = 'Arial'
    run.font.size = docx.shared.Pt(font_size)
    run.font.bold = bold

def populate_quotation(template_path, output_docx_path, data):
    doc = docx.Document(template_path)
    
    # --- Table 0: Customer & Job Info ---
    t0 = doc.tables[0]
    
    customer_name = sanitize_str(data.get('customerName', ''))
    p00 = t0.rows[0].cells[0].paragraphs[0]
    replace_text_in_paragraph(p00, '<Customer Name>', customer_name)
    
    quote_ref = sanitize_str(data.get('quoteReference', ''))
    p01 = t0.rows[0].cells[1].paragraphs[0]
    replace_text_in_paragraph(p01, '<Reference>', quote_ref)
    
    site_address = sanitize_str(data.get('siteAddress', ''))
    p10_cell = t0.rows[1].cells[0]
    for p in p10_cell.paragraphs:
        replace_text_in_paragraph(p, '<Site Address>', site_address)
        
    issued_by = sanitize_str(data.get('issuedBy', 'Prime Energy UK')) or 'Prime Energy UK'
    date_issued = sanitize_str(data.get('dateIssued', ''))
    valid_until = sanitize_str(data.get('validUntil', ''))
    
    p20 = t0.rows[2].cells[0].paragraphs[0]
    p20.text = f"Quotation issued by: {issued_by}"
    if len(p20.runs) > 0:
        p20.runs[0].font.name = 'Arial'
        p20.runs[0].font.size = docx.shared.Pt(9.5)
    
    p21 = t0.rows[2].cells[1].paragraphs[0]
    p21.text = f"Date issued: {date_issued}"
    if len(p21.runs) > 0:
        p21.runs[0].font.name = 'Arial'
        p21.runs[0].font.size = docx.shared.Pt(9.5)
    
    p22 = t0.rows[2].cells[2].paragraphs[0]
    p22.text = f"Valid until: {valid_until}"
    if len(p22.runs) > 0:
        p22.runs[0].font.name = 'Arial'
        p22.runs[0].font.size = docx.shared.Pt(9.5)

    # --- Table 1: Goods, Services, Commercial Totals ---
    t1 = doc.tables[1]
    
    goods_items = data.get('goodsItems', [])
    for i in range(4):
        r_idx = 1 + i
        row = t1.rows[r_idx]
        if i < len(goods_items):
            item = goods_items[i]
            desc = sanitize_str(item.get('description', ''))
            qty = item.get('quantity', 1)
            unit_price = item.get('unitPriceExVat', 0)
            total_price = item.get('totalPriceExVat', 0)
            set_cell_text(row.cells[0], desc)
            set_cell_text(row.cells[1], qty, align="center")
            set_cell_text(row.cells[2], format_currency(unit_price), align="right")
            set_cell_text(row.cells[3], format_currency(total_price), align="right")
        else:
            set_cell_text(row.cells[0], "")
            set_cell_text(row.cells[1], "")
            set_cell_text(row.cells[2], "")
            set_cell_text(row.cells[3], "")

    set_cell_text(t1.rows[5].cells[3], format_currency(data.get('goodsTotal', 0)), align="right", bold=True)
    
    services_items = data.get('servicesItems', [])
    for i in range(4):
        r_idx = 7 + i
        row = t1.rows[r_idx]
        if i < len(services_items):
            item = services_items[i]
            desc = sanitize_str(item.get('description', ''))
            qty = item.get('quantity', 1)
            unit_price = item.get('unitPriceExVat', 0)
            total_price = item.get('totalPriceExVat', 0)
            set_cell_text(row.cells[0], desc)
            set_cell_text(row.cells[1], qty, align="center")
            set_cell_text(row.cells[2], format_currency(unit_price), align="right")
            set_cell_text(row.cells[3], format_currency(total_price), align="right")
        else:
            set_cell_text(row.cells[0], "")
            set_cell_text(row.cells[1], "")
            set_cell_text(row.cells[2], "")
            set_cell_text(row.cells[3], "")

    set_cell_text(t1.rows[11].cells[3], format_currency(data.get('servicesTotal', 0)), align="right", bold=True)
    
    set_cell_text(t1.rows[12].cells[3], format_currency(data.get('totalContractValueExVat', 0)), align="right", bold=True)
    set_cell_text(t1.rows[13].cells[3], format_currency(data.get('vat', 0)), align="right")
    
    bus_deduction = data.get('busDeduction', 0)
    bus_str = f"-£{float(bus_deduction):,.2f}" if float(bus_deduction or 0) > 0 else "£0.00"
    set_cell_text(t1.rows[14].cells[3], bus_str, align="right", bold=True)
    
    set_cell_text(t1.rows[15].cells[3], format_currency(data.get('customerContributionIncVat', 0)), align="right", bold=True)

    # --- Table 2: Warranties ---
    t2 = doc.tables[2]
    ashp_brand = sanitize_str(data.get('ashpBrand', '')) or 'Daikin'
    cylinder_brand = sanitize_str(data.get('cylinderBrand', '')) or 'Joule'
    
    warranties = data.get('warranties', [
        {'item': 'Air Source Heat Pump', 'brand': ashp_brand, 'years': '5 Years'},
        {'item': 'Hot Water Cylinder', 'brand': cylinder_brand, 'years': '25 Years'},
        {'item': 'Radiators & Emitters', 'brand': 'City Plumbing', 'years': '10 Years'},
        {'item': 'Accessories & Controls', 'brand': 'Diverse', 'years': '2 Years'},
        {'item': 'Installation & Workmanship', 'brand': 'Prime Energy UK', 'years': '2 Years'}
    ])
    for i in range(min(5, len(t2.rows) - 1)):
        row = t2.rows[1 + i]
        if i < len(warranties):
            w = warranties[i]
            set_cell_text(row.cells[0], sanitize_str(w.get('item', '')), font_size=8.5)
            set_cell_text(row.cells[1], sanitize_str(w.get('brand', '')), font_size=8.0)
            set_cell_text(row.cells[2], sanitize_str(w.get('years', '')), align="center", font_size=8.5)

    doc.save(output_docx_path)

def main():
    if len(sys.argv) > 1 and os.path.exists(sys.argv[1]):
        with open(sys.argv[1], 'r', encoding='utf-8') as f:
            data = json.load(f)
    else:
        raw_input = sys.stdin.read()
        if not raw_input.strip():
            print(json.dumps({'error': 'No JSON payload provided.'}))
            sys.exit(1)
        data = json.loads(raw_input)

    template_path = data.get('templatePath')
    output_docx_path = data.get('outputDocxPath')
    output_pdf_path = data.get('outputPdfPath')
    ps_script_path = data.get('psScriptPath')

    if not template_path or not output_docx_path or not output_pdf_path:
        print(json.dumps({'error': 'templatePath, outputDocxPath, and outputPdfPath are required.'}))
        sys.exit(1)

    os.makedirs(os.path.dirname(os.path.abspath(output_docx_path)), exist_ok=True)
    os.makedirs(os.path.dirname(os.path.abspath(output_pdf_path)), exist_ok=True)

    # 1. Populate DOCX
    populate_quotation(template_path, output_docx_path, data)

    # 2. Convert to PDF using PowerShell MS Word COM
    if not ps_script_path:
        script_dir = os.path.dirname(os.path.abspath(__file__))
        ps_script_path = os.path.join(script_dir, 'convert_doc_to_pdf.ps1')

    cmd = [
        'powershell', '-ExecutionPolicy', 'Bypass', '-File', ps_script_path,
        '-docxPath', os.path.abspath(output_docx_path),
        '-pdfPath', os.path.abspath(output_pdf_path)
    ]

    res = subprocess.run(cmd, capture_output=True, text=True)
    if res.returncode != 0:
        print(json.dumps({
            'error': f"PDF conversion failed: {res.stderr or res.stdout}",
            'docxPath': output_docx_path
        }))
        sys.exit(1)

    print(json.dumps({
        'success': True,
        'docxPath': output_docx_path,
        'pdfPath': output_pdf_path
    }))

if __name__ == '__main__':
    main()
